"""Local vocabulary interchange and validated, reversible backup restore."""
import csv
import io
import json
import re
import shutil
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path, PurePosixPath

from utils.file_manager import repair_word_data, save_words, update_english_library

FIELDS = ['word', 'phonetic', 'part_of_speech', 'translation', 'example', 'analysis',
          'source_text', 'imported_at', 'generation_error', 'generation_mode']
ALIASES = {'单词': 'word', '词汇': 'word', '搭配': 'word', '英文': 'word', 'phrase': 'word',
           '音标': 'phonetic', '词性': 'part_of_speech', '释义': 'translation', '中文': 'translation',
           '例句': 'example', '解析': 'analysis', '原文': 'source_text', '导入时间': 'imported_at'}
MAX_SIZE = 512 * 1024 * 1024


def _rows_to_entries(rows):
    rows = [list(row) for row in rows if any(value is not None and str(value).strip() for value in row)]
    if not rows:
        return []
    header = [str(x or '').strip().lower() for x in rows[0]]
    header = [ALIASES.get(x, x) for x in header]
    if 'word' in header:
        if len([x for x in header if x in FIELDS]) != len(set(x for x in header if x in FIELDS)):
            raise ValueError('表头有重复字段，请保留每个字段一列。')
        data = [dict((key, '' if value is None else str(value)) for key, value in zip(header, row) if key in FIELDS)
                for row in rows[1:]]
    elif all(len(row) == 1 for row in rows):
        data = [{'word': str(row[0])} for row in rows]
    else:
        raise ValueError('多列表格必须包含 word（或“英文／词汇／搭配”）表头。')
    result, seen = [], set()
    for row in data:
        if not row.get('word', '').strip():
            raise ValueError('发现空的 word 单元格，请补全或删除该行。')
        row['word'] = re.sub(r'\s+', ' ', row['word']).strip()
        item = repair_word_data(row)
        if item['word'] in seen:
            raise ValueError(f"文件内存在重复搭配：{item['word']}，请先合并重复行。")
        seen.add(item['word'])
        result.append(item)
    return result


def read_exchange(path: Path):
    suffix = path.suffix.lower()
    if path.stat().st_size > 50 * 1024 * 1024:
        raise ValueError('导入文件过大（上限 50 MB）。')
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        with zipfile.ZipFile(path) as package:
            if sum(info.file_size for info in package.infolist()) > MAX_SIZE:
                raise ValueError('Excel 解压内容过大，请拆分后导入。')
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            sheet = book.active
            if sheet.max_row > 50000 or sheet.max_column > 100:
                raise ValueError('表格超过 50,000 行或 100 列，请拆分后导入。')
            rows = []
            for row in sheet.iter_rows():
                if any(cell.data_type == 'f' for cell in row):
                    raise ValueError('请将 Excel 公式转换为值后导入。')
                rows.append([cell.value for cell in row])
            return _rows_to_entries(rows)
        finally:
            book.close()
    raw = path.read_bytes()
    try:
        text = raw.decode('utf-8-sig')
    except UnicodeDecodeError:
        text = raw.decode('gb18030')
    if suffix == '.txt':
        # One entry per line; internal commas belong to the phrase.
        return _rows_to_entries([['word']] + [[line] for line in text.splitlines() if line.strip()])
    if suffix == '.csv':
        rows = [[value[1:] if len(value) > 1 and value[0] == "'" and value[1] in "'=+-@\t\r" else value
                 for value in row] for row in csv.reader(io.StringIO(text))]
        return _rows_to_entries(rows)
    raise ValueError('支持 .xlsx、.csv 和 .txt；旧版 .xls 请另存为 .xlsx。')


def export_exchange(path: Path, entries):
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.suffix.lower() == '.txt':
        path.write_text(''.join(x['word'] + '\n' for x in entries), encoding='utf-8')
    elif path.suffix.lower() == '.csv':
        with path.open('w', encoding='utf-8-sig', newline='') as output:
            writer = csv.DictWriter(output, fieldnames=FIELDS, extrasaction='ignore')
            writer.writeheader()
            for entry in entries:
                row = {key: str(entry.get(key, '')) for key in FIELDS}
                # Excel must not execute user content as formulas when opening CSV.
                row = {key: "'" + value if value and value[0] in "'=+-@\t\r" else value for key, value in row.items()}
                writer.writerow(row)
    elif path.suffix.lower() == '.xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        book = Workbook()
        sheet = book.active
        sheet.title = 'Vocabulary'
        sheet.append(FIELDS)
        for entry in entries:
            values = [str(entry.get(field, '')) for field in FIELDS]
            if any(len(value) > 32767 for value in values):
                raise ValueError('字段超过 Excel 单元格长度上限，请改用 CSV 或完整备份。')
            sheet.append(values)
        for row in sheet:
            for cell in row:
                if len(str(cell.value or '')) > 32767:
                    raise ValueError('字段超过 Excel 单元格长度上限，请改用 CSV 或完整备份。')
                cell.data_type = 's'  # Treat content as text, never as formulas.
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='334155')
        widths = [48, 28, 22, 36, 70, 45, 70, 23, 36, 20]
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        book.save(path)
        book.close()
    else:
        raise ValueError('请选择 .xlsx、.csv 或 .txt 格式。')


def merge_entries(current, incoming, replace=False):
    merged = {x['word']: dict(x) for x in current}
    for item in incoming:
        old = merged.get(item['word'])
        if old is not None and not replace:
            continue
        result = dict(item)
        result['imported_at'] = result.get('imported_at') or (old or {}).get('imported_at') or datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        merged[item['word']] = result
    return list(merged.values())


def create_backup(root: Path, output: Path):
    root, output = root.resolve(), output.resolve()
    if output.is_relative_to(root / 'data') or output.is_relative_to(root / 'audio'):
        raise ValueError('请将备份保存到 data 和 audio 目录之外。')
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_suffix(output.suffix + '.tmp')
    try:
        with zipfile.ZipFile(temporary, 'w', zipfile.ZIP_DEFLATED) as archive:
            archive.writestr('manifest.json', json.dumps({'format': 'AnkiGen', 'version': 1}))
            archive.write(root / 'data' / 'words.json', 'data/words.json')
            articles = root / 'data' / 'articles.json'
            if articles.exists():
                archive.write(articles, 'data/articles.json')
            else:
                archive.writestr('data/articles.json', '[]')
            for audio in sorted((root / 'audio').glob('*.mp3')):
                archive.write(audio, 'audio/' + audio.name)
        inspect_backup(temporary)
        temporary.replace(output)
    finally:
        temporary.unlink(missing_ok=True)


def inspect_backup(path: Path):
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        if len(infos) > 20000 or sum(info.file_size for info in infos) > MAX_SIZE:
            raise ValueError('备份内容超过安全读取上限（512 MB / 20,000 个文件）。')
        names = set()
        for info in infos:
            name = info.filename
            parts = PurePosixPath(name).parts
            valid = name in {'manifest.json', 'data/words.json', 'data/articles.json'} or (
                len(parts) == 2 and parts[0] == 'audio' and parts[1].endswith('.mp3')
                and not any(char in parts[1] for char in '\\:') and parts[1] not in {'.', '..'}
            )
            if not valid or name.casefold() in names or info.is_dir():
                raise ValueError(f'备份含不支持或重复的文件路径：{name}')
            names.add(name.casefold())
        manifest = json.loads(archive.read('manifest.json'))
        if manifest != {'format': 'AnkiGen', 'version': 1}:
            raise ValueError('不是受支持的 AnkiGen 备份版本。')
        entries = json.loads(archive.read('data/words.json'))
        articles = json.loads(archive.read('data/articles.json'))
        if not isinstance(entries, list) or not isinstance(articles, list):
            raise ValueError('词库或文章数据格式无效。')
        for item in entries:
            if not isinstance(item, dict) or not isinstance(item.get('word'), str) or not item['word'].strip():
                raise ValueError('备份含无效词条。')
            if any(not isinstance(value, str) for value in item.values()):
                raise ValueError('词条字段必须是文本。')
        words = [item['word'].strip().lower() for item in entries]
        if len(set(words)) != len(words):
            raise ValueError('备份含重复词条。')
        if any(not isinstance(row, dict) or not isinstance(row.get('text'), str) for row in articles):
            raise ValueError('文章数据格式无效。')
        return {'entries': entries, 'articles': articles, 'audio_count': sum(n.startswith('audio/') for n in names)}


def restore_backup(root: Path, backup: Path):
    root, backup = root.resolve(), backup.resolve()
    data = inspect_backup(backup)
    recovery = root / 'backups' / ('before-restore-' + datetime.now().strftime('%Y%m%d-%H%M%S-%f') + '.zip')
    create_backup(root, recovery)
    with tempfile.TemporaryDirectory(prefix='.restore-', dir=root) as directory:
        stage = Path(directory).resolve()
        assert stage.parent == root
        staged_data, staged_audio = stage / 'data', stage / 'audio'
        # Preserve this computer's settings/API keys; backups never contain them.
        shutil.copytree(root / 'data', staged_data)
        staged_audio.mkdir()
        (staged_data / 'words.json').write_text(json.dumps(data['entries'], ensure_ascii=False, indent=2), encoding='utf-8')
        (staged_data / 'articles.json').write_text(json.dumps(data['articles'], ensure_ascii=False, indent=2), encoding='utf-8')
        with zipfile.ZipFile(backup) as archive:
            for name in archive.namelist():
                if name.startswith('audio/'):
                    target = staged_audio / PurePosixPath(name).name
                    assert target.resolve().parent == staged_audio.resolve()
                    target.write_bytes(archive.read(name))
        swapped = []
        try:
            for name in ('data', 'audio'):
                target, old = root / name, stage / ('old-' + name)
                assert target.resolve().parent == root and old.resolve().parent == stage
                target.rename(old)
                try:
                    (stage / name).rename(target)
                except Exception:
                    old.rename(target)
                    raise
                swapped.append(name)
            update_english_library(root / 'data' / 'words.json', data['entries'])
        except Exception:
            for name in reversed(swapped):
                (root / name).rename(stage / ('failed-' + name))
                (stage / ('old-' + name)).rename(root / name)
            raise
    return recovery
